# -*- coding: utf-8 -*-
"""Выгрузки xlsx -> dist/data.json (все срезы дашборда)."""
import json, re, collections
from urllib.parse import urlsplit
import openpyxl
import config as C

URL_RE = re.compile(r"https?://[^\s<>\"'\]\)（）]+")
SERVICE = [re.compile(p, re.I) for p in C.SERVICE_URL_PATTERNS]


def read_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(v is not None for v in r)]


def pick(row, *names):
    for n in names:
        if n in row and row[n] is not None:
            return row[n]
    return None


def clean_url(u):
    u = u.rstrip(".,;:!?*»)]}>'\"")
    return u


def host_of(u):
    try:
        h = urlsplit(u).hostname or ""
    except ValueError:
        return ""
    h = h.lower().strip(".")
    if h.startswith("www."):
        h = h[4:]
    return h


def extract_urls(sources, text):
    out, seen = [], set()
    for chunk in (sources or "", text or ""):
        for u in URL_RE.findall(str(chunk)):
            u = clean_url(u)
            if any(p.search(u) for p in SERVICE):
                continue
            h = host_of(u)
            if not h or "." not in h:
                continue
            key = u.split("#")[0]
            key = re.sub(r"[?&]utm_[^&]+", "", key)
            if key in seen:
                continue
            seen.add(key)
            out.append((u, h, urlsplit(u).path or "/"))
    return out


def strip_urls(text):
    t = re.sub(r"\]\((https?://[^)]+)\)", "]", str(text or ""))
    return URL_RE.sub(" ", t)


BRAND_RE = {}
for k, b in C.BRANDS.items():
    flags = 0 if b.get("cs") else re.I
    BRAND_RE[k] = [re.compile(p, flags) for p in b["patterns"]]


def brand_positions(text):
    pos = {}
    for k, rs in BRAND_RE.items():
        best = None
        for r in rs:
            m = r.search(text)
            if m and (best is None or m.start() < best):
                best = m.start()
        if best is not None:
            pos[k] = best
    return pos


def host_match(h, domain):
    return h == domain or h.endswith("." + domain)


def load_answers():
    ans = []
    for src in ("models", "google", "alice"):
        for r in read_sheet(C.INPUT[src]):
            ch_raw = str(pick(r, "AI-модель", "Модель") or "").strip().lower()
            ch = C.CHANNEL_ALIASES.get(ch_raw)
            if not ch:
                continue
            text = pick(r, "Ответ", "Нейроответ") or ""
            src_field = pick(r, "Источники") or ""
            ans.append({
                "ch": ch, "type": C.CHANNELS[ch][1],
                "niche": str(pick(r, "Ниша") or "").strip(),
                "q": str(pick(r, "Запрос") or "").strip(),
                "text": str(text),
                "urls": extract_urls(src_field, text),
                "pos": brand_positions(strip_urls(text)),
            })
    return ans


def pct(a, b):
    return None if not b else round(100.0 * a / b, 1)


def slice_metrics(rows, brand_keys, target, cit_ok, hub=None):
    """rows — ответы среза; cit_ok — множество каналов с доступными источниками."""
    n = len(rows)
    cit_rows = [r for r in rows if r["ch"] in cit_ok]
    nc = len(cit_rows)
    res = {"n": n, "nc": nc, "brands": {}}
    total_mentions = 0
    none_cnt = 0
    stats = {k: {"m": 0, "ranks": [], "first": 0, "cit": 0, "hosts": collections.Counter()} for k in brand_keys}
    hub_n = 0
    for r in rows:
        present = sorted([k for k in brand_keys if k in r["pos"]], key=lambda k: r["pos"][k])
        if not present:
            none_cnt += 1
        total_mentions += len(present)
        for i, k in enumerate(present):
            stats[k]["m"] += 1
            stats[k]["ranks"].append(i + 1)
            if i == 0:
                stats[k]["first"] += 1
    for r in cit_rows:
        hosts = {h for _, h, _ in r["urls"]}
        for k in brand_keys:
            own = {h for h in hosts if any(host_match(h, d) for d in C.BRANDS[k]["domains"])}
            if own:
                stats[k]["cit"] += 1
                for h in own:
                    stats[k]["hosts"][h] += 1
        if hub and any((h == hub[0]) and p.startswith(hub[1]) for _, h, p in r["urls"]):
            hub_n += 1
    for k, s in stats.items():
        res["brands"][k] = {
            "vis": pct(s["m"], n), "cnt": s["m"],
            "sov": pct(s["m"], total_mentions),
            "pos": round(sum(s["ranks"]) / len(s["ranks"]), 2) if s["ranks"] else None,
            "first": pct(s["first"], s["m"]),
            "cit": pct(s["cit"], nc), "citn": s["cit"],
            # разбивка цитирования по сайтам бренда (один ответ может сослаться на несколько)
            "hosts": [{"h": h, "n": c, "p": pct(c, nc)} for h, c in s["hosts"].most_common()],
        }
    if hub:
        res["hub"] = {"n": hub_n, "p": pct(hub_n, nc)}
    res["none"] = pct(none_cnt, n)
    # ТОП доменов: доля ответов (из цитирующих каналов), где домен указан в источниках
    dc = collections.Counter()
    for r in cit_rows:
        for h in {h for _, h, _ in r["urls"]}:
            dc[h] += 1
    top = []
    for h, c in dc.most_common(C.TOP_DOMAINS_N):
        owner = next((k for k in C.BRANDS if any(host_match(h, d) for d in C.BRANDS[k]["domains"])), None)
        top.append({"d": h, "n": c, "p": pct(c, nc), "b": owner if owner in brand_keys else (owner and "other_brand")})
    res["top"] = top
    res["domains_total"] = len(dc)
    return res


def gaps(rows, brand_keys):
    """Запросы, где конкуренты названы, а ОН КЛИНИК — нет (по всем каналам среза)."""
    byq = collections.defaultdict(lambda: {"ch": set(), "comp": collections.Counter(), "client": 0, "n": 0})
    for r in rows:
        g = byq[r["q"]]
        g["n"] += 1
        if "onclinic" in r["pos"]:
            g["client"] += 1
            continue
        comps = [k for k in brand_keys if k != "onclinic" and k in r["pos"]]
        if comps:
            g["ch"].add(r["ch"])
            for k in comps:
                g["comp"][k] += 1
    out = []
    for q, g in byq.items():
        if g["client"] == 0 and g["comp"]:
            out.append({"q": q, "n": g["n"], "ch": sorted(g["ch"], key=list(C.CHANNELS).index),
                        "comp": [k for k, _ in g["comp"].most_common()], "w": sum(g["comp"].values())})
    out.sort(key=lambda x: (-len(x["ch"]), -x["w"]))
    return out[: C.GAP_PROMPTS_N], len(out)


def metrika():
    wb = openpyxl.load_workbook(C.INPUT["metrika"], data_only=True)
    ws = wb.worksheets[0]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    title = rows[0][0] if rows and rows[0] else ""
    models, types, mode = [], [], None
    for r in rows:
        a = str(r[0]).strip() if r and r[0] is not None else ""
        if a == "Модель":
            mode = "m"; continue
        if a == "Тип трафика":
            mode = "t"; continue
        if not a or a.startswith("Итого") or a in ("Разбивка по моделям", "Информационный и коммерческий трафик"):
            if a.startswith("Итого"):
                mode = None
            continue
        if mode == "m":
            label = C.METRIKA_MODELS.get(a.lower())
            if label:  # сервисы вне замера не показываем
                models.append({"name": label, "visits": int(float(r[1] or 0)), "goals": int(float(r[2] or 0))})
        elif mode == "t":
            types.append({"name": a, "visits": int(float(r[1] or 0))})
    period = re.search(r"(\d{2}\.\d{2}\.\d{4})\s*[–-]\s*(\d{2}\.\d{2}\.\d{4})", str(title))
    return {"title": title, "period": period.groups() if period else None, "models": models, "types": types}


def main():
    ans = load_answers()
    # доступность источников по каналам
    avail = {}
    for ch in C.CHANNELS:
        rs = [a for a in ans if a["ch"] == ch]
        avail[ch] = pct(sum(1 for a in rs if a["urls"]), len(rs)) if rs else None
    cit_ok = {ch for ch, v in avail.items() if v is not None and v >= C.CITATION_MIN_SHARE * 100}

    scopes = C.DIRECTIONS + C.FAQ_GROUPS
    data = {"scopes": {}, "channels": {k: {"name": v[0], "type": v[1], "icon": v[2], "src": avail[k], "cit": k in cit_ok}
                                        for k, v in C.CHANNELS.items()}}
    ch_keys = ["models", "search"] + list(C.CHANNELS)
    for s in scopes:
        rows = [a for a in ans if a["niche"] == s["niche"]]
        bk = C.COMPETITOR_SETS[s["competitors"]]
        sd = {"title": s["title"], "brands": bk, "target_label": s["target_label"],
              "hub_label": (s["hub"][0] + s["hub"][1]) if s.get("hub") else None,
              "prompts": {t: len({a["q"] for a in rows if a["type"] == t}) for t in ("models", "search")},
              "cube": {}, "gaps": {}}
        for ck in ch_keys:
            sub = [a for a in rows if (a["type"] == ck if ck in ("models", "search") else a["ch"] == ck)]
            if not sub:
                continue
            sd["cube"][ck] = slice_metrics(sub, bk, s["target"], cit_ok, s.get("hub"))
        for t in ("models", "search"):
            g, total = gaps([a for a in rows if a["type"] == t], bk)
            sd["gaps"][t] = {"list": g, "total": total}
        data["scopes"][s["key"]] = sd

    data["brands"] = {k: {"name": b["name"], "client": b.get("client", False), "domain": b["domains"][0]} for k, b in C.BRANDS.items()}
    data["directions"] = [{"key": s["key"], "title": s["title"]} for s in C.DIRECTIONS]
    data["faq"] = [{"key": s["key"], "title": s["title"]} for s in C.FAQ_GROUPS]
    data["totals"] = {
        "answers": len(ans),
        "prompts_models": len({(a["niche"], a["q"]) for a in ans if a["type"] == "models"}),
        "prompts_search": len({(a["niche"], a["q"]) for a in ans if a["type"] == "search"}),
        "answers_models": sum(1 for a in ans if a["type"] == "models"),
        "answers_search": sum(1 for a in ans if a["type"] == "search"),
        "citations": sum(len(a["urls"]) for a in ans if a["ch"] in cit_ok),
        "domains": len({h for a in ans if a["ch"] in cit_ok for _, h, _ in a["urls"]}),
    }
    data["metrika"] = metrika()
    data["meta"] = {"date": C.MEASURE_DATE, "client": C.CLIENT, "site": C.CLIENT_SITE, "cit_min": C.CITATION_MIN_SHARE}
    with open("dist/data.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print("answers", len(ans), "avail", avail, "totals", data["totals"])


if __name__ == "__main__":
    main()
